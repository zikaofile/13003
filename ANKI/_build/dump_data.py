# -*- coding: utf-8 -*-
"""读取 13003大纲考点.xlsx Sheet1 全量数据并落盘 JSON（供后续卡片生成）。"""
import json
import sys
from openpyxl import load_workbook

SRC = r"F:\syncthing\考试\自考\13003数据结构与算法\13003大纲考点.xlsx"
OUT = r"F:\syncthing\考试\自考\13003数据结构与算法\ANKI\_build\rows.json"

wb = load_workbook(SRC, data_only=True, read_only=True)
ws = wb["Sheet1"]

rows = []
for i, row in enumerate(ws.iter_rows(min_row=3, max_col=7, values_only=True), start=3):
    rows.append({
        "excel_row": i,
        "chapter": (row[0] if row[0] is not None else ""),
        "requirement": (row[1] if row[1] is not None else ""),
        "summary": (row[2] if row[2] is not None else ""),
        "knowledge": (row[3] if row[3] is not None else ""),
        "desc": (row[4] if row[4] is not None else ""),
        "seq": row[5],
        "page": row[6],
    })
wb.close()

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False, indent=1)

# 概要统计
total = len(rows)
has_k = sum(1 for r in rows if r["knowledge"].strip())
has_e = sum(1 for r in rows if r["desc"].strip())
has_page = sum(1 for r in rows if r["page"] is not None)
chapters = []
for r in rows:
    ch = r["chapter"].strip()
    if ch and ch not in chapters:
        chapters.append(ch)
reqs = []
for r in rows:
    q = r["requirement"].strip()
    if q and q not in reqs:
        reqs.append(q)
summaries = []
for r in rows:
    s = r["summary"].strip()
    if s and s not in summaries:
        summaries.append(s)
print("总行数:", total)
print("含知识点D:", has_k, "含描述E:", has_e, "含页码G:", has_page)
print("章节数:", len(chapters), chapters)
print("要求B取值:", reqs)
print("小结C数量:", len(summaries))
print("小结C取值:", summaries)
